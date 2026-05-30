"""
PASSO 2: Tratamento dos CSVs brutos de tramitação.

Para cada CSV de comissão (em relatorios_brutos/):
  - Divide "Tramitação" -> "Data Tramitação" + "Local de Tramitação"
  - Divide "Tramitação Interna" -> "Data da Tramitação Interna"
                                 + "Local de Tramitação Interna"
                                 + "Motivo Tramitação Interna"
  - Salva CSV tratado em relatorios_tratados/
  - Consolida num único Excel com 7 abas (uma por comissão)

Dependências:
  pip install pandas openpyxl
"""

import os
import re
from pathlib import Path
import pandas as pd

COMISSOES = ["ADM", "CCJ", "ECON", "EDUC", "FIN", "SAUDE", "URB"]
DIR_BRUTOS = Path(__file__).parent / "relatorios_brutos"
DIR_TRATADOS = Path(__file__).parent / "relatorios_tratados"
EXCEL_SAIDA = DIR_TRATADOS / "acervo_comissoes_tratado.xlsx"

RE_DATA = re.compile(r"(\d{1,2}/\d{1,2}/\d{2,4})")


# ---------------------------------------------------------------------------
# Carregamento
# ---------------------------------------------------------------------------

def detectar_separador(caminho: Path) -> str:
    texto = caminho.read_text(encoding="utf-8-sig", errors="replace")
    primeira = texto.split("\n")[0]
    for sep in [";", ",", "\t"]:
        if primeira.count(sep) > 0:
            return sep
    return ";"


def carregar_csv(caminho: Path) -> pd.DataFrame:
    sep = detectar_separador(caminho)
    for enc in ("utf-8-sig", "latin-1", "cp1252"):
        try:
            df = pd.read_csv(caminho, sep=sep, encoding=enc, dtype=str, on_bad_lines="skip")
            df.columns = [c.strip() for c in df.columns]
            return df
        except Exception:
            continue
    raise RuntimeError(f"Não foi possível ler o arquivo: {caminho}")


def encontrar_coluna(df: pd.DataFrame, candidatos: list) -> str | None:
    """Retorna o nome real da coluna que corresponde a um dos candidatos (busca flexível)."""
    cols_norm = {re.sub(r"[^a-z]", "", c.lower()): c for c in df.columns}
    for cand in candidatos:
        chave = re.sub(r"[^a-z]", "", cand.lower())
        if chave in cols_norm:
            return cols_norm[chave]
    # Busca parcial
    for cand in candidatos:
        chave = re.sub(r"[^a-z]", "", cand.lower())
        for col_norm, col_orig in cols_norm.items():
            if chave in col_norm:
                return col_orig
    return None


# ---------------------------------------------------------------------------
# Divisão das colunas compostas
# ---------------------------------------------------------------------------

def dividir_tramitacao(valor) -> pd.Series:
    """
    Divide 'Tramitação' (formato: 'dd/mm/aaaa Local') em:
      Data Tramitação | Local de Tramitação
    """
    vazio = pd.Series({"Data Tramitação": "", "Local de Tramitação": ""})
    if pd.isna(valor) or not str(valor).strip():
        return vazio

    texto = str(valor).strip()
    m = RE_DATA.search(texto)
    if m:
        data = m.group(1)
        local = texto[m.end():].strip().lstrip("-–").strip()
    else:
        data = ""
        local = texto

    return pd.Series({"Data Tramitação": data, "Local de Tramitação": local})


def dividir_tramitacao_interna(valor) -> pd.Series:
    """
    Divide 'Tramitação Interna' (formato: 'dd/mm/aaaa Local - Motivo') em:
      Data da Tramitação Interna | Local de Tramitação Interna | Motivo Tramitação Interna
    """
    vazio = pd.Series({
        "Data da Tramitação Interna": "",
        "Local de Tramitação Interna": "",
        "Motivo Tramitação Interna": "",
    })
    if pd.isna(valor) or not str(valor).strip():
        return vazio

    texto = str(valor).strip()
    m = RE_DATA.search(texto)
    if m:
        data = m.group(1)
        resto = texto[m.end():].strip().lstrip("-–").strip()
    else:
        data = ""
        resto = texto

    # Separa local e motivo pelo primeiro " - " ou " – "
    sep = re.search(r"\s[-–]\s", resto)
    if sep:
        local = resto[: sep.start()].strip()
        motivo = resto[sep.end():].strip()
    else:
        local = resto
        motivo = ""

    return pd.Series({
        "Data da Tramitação Interna": data,
        "Local de Tramitação Interna": local,
        "Motivo Tramitação Interna": motivo,
    })


# ---------------------------------------------------------------------------
# Tratamento principal
# ---------------------------------------------------------------------------

def tratar_dataframe(df: pd.DataFrame, comissao: str) -> pd.DataFrame:
    col_tram = encontrar_coluna(df, ["Tramitação", "Tramitacao"])
    col_tram_int = encontrar_coluna(df, ["Tramitação Interna", "Tramitacao Interna"])

    if col_tram is None:
        print(f"  AVISO [{comissao}]: 'Tramitação' não encontrada. Colunas: {list(df.columns)}")
    if col_tram_int is None:
        print(f"  AVISO [{comissao}]: 'Tramitação Interna' não encontrada. Colunas: {list(df.columns)}")

    cols = list(df.columns)

    # Divide Tramitação
    if col_tram and col_tram in df.columns:
        idx = cols.index(col_tram)
        novas = df[col_tram].apply(dividir_tramitacao)
        df = df.drop(columns=[col_tram])
        for i, nova in enumerate(["Data Tramitação", "Local de Tramitação"]):
            df.insert(idx + i, nova, novas[nova])
        cols = list(df.columns)

    # Divide Tramitação Interna
    if col_tram_int and col_tram_int in df.columns:
        idx = cols.index(col_tram_int)
        novas = df[col_tram_int].apply(dividir_tramitacao_interna)
        df = df.drop(columns=[col_tram_int])
        for i, nova in enumerate([
            "Data da Tramitação Interna",
            "Local de Tramitação Interna",
            "Motivo Tramitação Interna",
        ]):
            df.insert(idx + i, nova, novas[nova])

    return df


# ---------------------------------------------------------------------------
# Excel com formatação básica
# ---------------------------------------------------------------------------

def escrever_aba_excel(writer, df: pd.DataFrame, nome_aba: str):
    df.to_excel(writer, sheet_name=nome_aba, index=False)
    ws = writer.sheets[nome_aba]

    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    # Cabeçalho em negrito com fundo azul claro
    header_fill = PatternFill("solid", fgColor="BDD7EE")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    # Ajusta largura das colunas
    for col_idx, col_cells in enumerate(ws.columns, 1):
        max_len = max(
            (len(str(c.value)) if c.value is not None else 0) for c in col_cells
        )
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 3, 50)

    # Congela a primeira linha
    ws.freeze_panes = "A2"


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    DIR_TRATADOS.mkdir(parents=True, exist_ok=True)

    dfs: dict[str, pd.DataFrame] = {}
    erros = []

    for comissao in COMISSOES:
        caminho = DIR_BRUTOS / f"acervo_{comissao}.csv"
        if not caminho.exists():
            print(f"[{comissao}] Arquivo não encontrado: {caminho}")
            erros.append(comissao)
            continue

        print(f"\n[{comissao}] Carregando {caminho.name}...")
        try:
            df = carregar_csv(caminho)
            print(f"  {len(df)} linhas | Colunas originais: {list(df.columns)}")

            df_tratado = tratar_dataframe(df, comissao)
            print(f"  Colunas após tratamento: {list(df_tratado.columns)}")

            saida_csv = DIR_TRATADOS / f"acervo_{comissao}_tratado.csv"
            df_tratado.to_csv(saida_csv, sep=";", index=False, encoding="utf-8-sig")
            print(f"  CSV salvo: {saida_csv.name}")

            dfs[comissao] = df_tratado

        except Exception as e:
            print(f"  ERRO [{comissao}]: {e}")
            erros.append(comissao)

    # Excel consolidado
    if dfs:
        print(f"\nGerando Excel: {EXCEL_SAIDA.name}")
        with pd.ExcelWriter(EXCEL_SAIDA, engine="openpyxl") as writer:
            for comissao in COMISSOES:
                if comissao in dfs:
                    escrever_aba_excel(writer, dfs[comissao], comissao)
                    print(f"  Aba '{comissao}': {len(dfs[comissao])} linhas")
        print(f"Excel salvo: {EXCEL_SAIDA}")

    print(f"\n{'='*50}")
    print("RESUMO PASSO 2")
    print(f"{'='*50}")
    print(f"Tratados: {list(dfs.keys())}")
    if erros:
        print(f"Não encontrados/erros: {erros}")


if __name__ == "__main__":
    main()
